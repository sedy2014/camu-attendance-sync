#!/usr/bin/env python3
"""
build_attendance_grid.py

Generic builder/updater for a camu.in-style weekly attendance grid workbook.
Institution-agnostic: takes a JSON spec (see examples/sample_config.json)
describing one or more "groups" (e.g. a semester, a batch, a program) and
writes/updates one sheet per group in a single .xlsx workbook.

Grid layout per sheet:
    Row 1: Title
    Row 3: Section | DAY1 | DAY2 | ... (yellow header)
    Row 4: (blank) | date1 | date2 | ...
    Row 5+: one row per section, one cell per day:
        "present/total"  -> confirmed attendance count
        "NA"             -> confirmed no class scheduled that day
        "?"              -> a class was held but the count hasn't been looked up yet
        ""                -> day not checked at all

Usage:
    python build_attendance_grid.py --config path/to/config.json

The config's "output_path" is the workbook to create/update. If it already
exists, existing sheets are preserved and only the sheets named in this
config are added or overwritten (so you can run this once per week without
losing prior weeks).
"""

import argparse
import json
import os
import sys

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Arial"
HEADER_FONT = Font(name=FONT_NAME, bold=True)
HEADER_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
BODY_FONT = Font(name=FONT_NAME)
NOTE_FONT = Font(name=FONT_NAME, italic=True, size=9, color="808080")
UNKNOWN_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FOOTER_NOTE = (
    "Source: camu.in Schedule -> class -> Attendance page (read via an already logged-in "
    "browser session; no credentials stored or entered automatically). "
    "NA = confirmed no class scheduled that day. ? (shaded) = a class was held but the count "
    "hasn't been looked up yet. Blank = day not checked at all."
)


def cell_text(value):
    """value is either [present, total], or the literal strings 'NA' / '?' / '' """
    if value is None:
        return ""
    if isinstance(value, (list, tuple)) and len(value) == 2:
        present, total = value
        return f"{present}/{total}"
    return str(value)


def build_sheet(wb, group):
    sheet_id = group["sheet_id"]
    week_label = group.get("week_label", "")
    sheet_name = f"{sheet_id}_{week_label}" if week_label else sheet_id
    sheet_name = sheet_name[:31]  # Excel sheet name limit

    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(title=sheet_name)

    title = group.get("title", sheet_id)
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT_NAME, bold=True, size=13)

    days = group["days"]
    day_dates = group.get("day_dates", [""] * len(days))
    sections = group["sections"]
    data = group.get("data", {})
    course_notes = group.get("course_notes", {})

    header_row = 3
    hc = ws.cell(row=header_row, column=1, value="Section")
    hc.font, hc.fill, hc.border = HEADER_FONT, HEADER_FILL, BORDER
    hc.alignment = Alignment(horizontal="center", vertical="center")

    for i, day in enumerate(days):
        col = i + 2
        c = ws.cell(row=header_row, column=col, value=str(day).upper())
        c.font, c.fill, c.border = HEADER_FONT, HEADER_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c2 = ws.cell(row=header_row + 1, column=col, value=day_dates[i] if i < len(day_dates) else "")
        c2.font = NOTE_FONT
        c2.alignment = Alignment(horizontal="center")

    ws.cell(row=header_row + 1, column=1, value="").border = BORDER

    r = header_row + 2
    for section in sections:
        cell = ws.cell(row=r, column=1, value=section)
        cell.font, cell.border = BODY_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")

        day_map = data.get(section, {})
        for i in range(len(days)):
            raw = day_map.get(str(i), day_map.get(i))
            text = cell_text(raw)
            c = ws.cell(row=r, column=i + 2, value=text)
            c.font, c.border = BODY_FONT, BORDER
            c.alignment = Alignment(horizontal="center", vertical="center")
            if text == "?":
                c.fill = UNKNOWN_FILL
        r += 1

    note_row = r + 1
    if course_notes:
        ws.cell(row=note_row, column=1, value="Course reference:").font = Font(
            name=FONT_NAME, bold=True, size=9
        )
        note_row += 1
        for section in sections:
            if section in course_notes:
                ws.cell(row=note_row, column=1, value=section).font = NOTE_FONT
                ws.cell(row=note_row, column=2, value=course_notes[section]).font = NOTE_FONT
                note_row += 1
        note_row += 1

    ws.cell(row=note_row, column=1, value=FOOTER_NOTE).font = NOTE_FONT

    widths = [10] + [12] * len(days)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", required=True, help="Path to a JSON config file")
    args = parser.parse_args()

    with open(args.config, "r", encoding="utf-8") as f:
        spec = json.load(f)

    output_path = spec["output_path"]
    groups = spec["groups"]

    if os.path.exists(output_path):
        wb = openpyxl.load_workbook(output_path)
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

    for group in groups:
        build_sheet(wb, group)

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    wb.save(output_path)
    print(f"Saved: {output_path}")
    print(f"Sheets: {wb.sheetnames}")


if __name__ == "__main__":
    sys.exit(main())
